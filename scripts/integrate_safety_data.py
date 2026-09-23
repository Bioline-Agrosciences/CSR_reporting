"""
Integrates Health & Safety data — from the two sources that own it — into the
consolidated raw data (input_data/Raw_data_CSR.xlsx), the same file
integrate_data_entry.py feeds for the CSR-referent-owned indicators.

Unlike the other indicators (Water/Energy/Refrigerants/Waste/ESG, entered
monthly by each BU's CSR referent), the Safety indicators are NOT owned by
the CSR referent (see Indicator_reference_CSR.xlsx, "Responsible" column:
H&S manager / HR services) and are not part of the <BU>_data_entry_CSR_<year>.xlsm
file generate_data_entry_file.py produces. They come from two other systems:

  - config.WORKING_HOURS_FILE  -> Saf.4.2 "Working hours of the month"
    One fixed file, wide format (1 row per BU, 1 column per month), updated
    in place. A blank cell means the month hasn't been declared yet (left
    out here, not defaulted to 0 or 0 hours). No fallback exists anymore if
    Saf.4.2 is missing (there used to be one, based on Env.1/FTE — both
    gone from this pipeline as of 23/09/2026, see csr_calc_engine.py):
    Saf.6/Saf.7 simply can't be computed downstream for a month with no
    Saf.4.2, same as any other missing raw component.

  - config.ACCIDENTS_DIR / config.ACCIDENTS_FILE_PATTERN -> the BlueKanGo
    accidents export, one row per accident/incident. Picks the most recently
    modified file matching the pattern (a new export is a new file, its name
    carries an export timestamp). From it we derive, per BU and month:
      * Saf.1 "Number of days without an accident" — see
        compute_days_without_accident() below.
      * Saf.2 "Non lost time injury"  = count of accidents with no work
        stoppage ("Arrêt de travail" = NON)
      * Saf.3 "Lost time injury"      = count of accidents WITH a work
        stoppage ("Arrêt de travail" = OUI)
      * Saf.5 "Days lost"             = sum of "Durée totale de l'arrêt"
        for accidents with a work stoppage
    Two filters are applied everywhere in this module, matching the KPI
    definitions in the reference list ("Number of accidents (Bioline
    employees)..."):
      * accidents flagged "Intérimaire / Agency / non Bioline employees" =
        Oui are EXCLUDED (temp/agency workers are not Bioline employees) —
        applied to Saf.1 too, for consistency with the rest of the H&S
        indicators (an assumption — flag it to Aurélie if agency-worker
        accidents should still reset the "days without accident" counter);
      * Saf.2/Saf.3/Saf.5 additionally only count accidents dated in
        TARGET_YEAR (they are monthly aggregates); Saf.1 deliberately does
        NOT restrict to TARGET_YEAR — see compute_days_without_accident.
    A BU/month with zero qualifying accidents never appears in the raw
    aggregation — fill_zero_accident_months() turns that absence into an
    explicit 0 for Saf.2/Saf.3/Saf.5, for every (BU, Month) that IS covered
    by the working hours file (our proxy for "this month is active for
    reporting") — "no accident" is a real, reportable value, not missing data.

Once these rows are in Raw_data_CSR.xlsx, Saf.4.2/Saf.2/Saf.3/Saf.5 are
available for Saf.6 (Frequency rate) and Saf.7 (Gravity rate) to be
computed downstream (Power BI/Fabric) — NOT by csr_calc_engine.py: as of
23/09/2026 that script only computes additive (sum) indicators, never
ratios, because a ratio can't be validly rolled up to a group total the
way a sum can (see its module docstring). Saf.6/Saf.7 are still classified
"calculated" in the reference list, just no longer produced by that
script's FORMULAS dict.

NOT handled here (deliberately out of scope, see the analysis shared with
Aurélie):
  - Env.1 "Sales" (SAP) — this pipeline has no real sales/SAP figures at
    all; Env.1 is excluded from the reference list and from
    Raw_data_CSR.xlsx entirely (see extract_reference_and_data.py's
    EXCLUDED_IDS). The Sales/SAP join happens downstream, in Fabric.
  - Wat.2, Ene.10, Ene.11, Was.3, Was.4 (ratios, same reasoning as Saf.6/
    Saf.7 above) — also computed downstream, not here.

Usage
-----
    uv run scripts/integrate_safety_data.py

Follow up with:
    uv run scripts/csr_calc_engine.py

Dependencies: pandas, openpyxl.
"""
import calendar
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config
from integrate_data_entry import load_consolidated, export_consolidated

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

# BlueKanGo's free-text "Business Unit" label -> our internal BU code.
# "Bioline Viridaxis" is a guess (no accident for that BU in the sample
# export, so this mapping is unverified) — to confirm with Aurélie. Any
# label NOT in this dict is reported as a warning (not silently dropped),
# same philosophy as the "NOT PARSABLE" data quality notes elsewhere in the
# pipeline: surface it rather than guess.
BU_NAME_MAP = {
    "Bioline Africa": "BAF",
    "Bioline US": "BUS",
    "Bioline UK": "BUK",
    "Bioline Iberia": "BIB",
    "Bioline France": "BFR",
    "Bioline Viridaxis": "Viridaxis",  # unverified — no example row seen yet
}


def _find_latest_accidents_file() -> Path:
    """Returns the most recently modified file matching
    config.ACCIDENTS_FILE_PATTERN in config.ACCIDENTS_DIR — the BlueKanGo
    export file name carries an export timestamp, so it's a different file
    every time, never a fixed name."""
    candidates = sorted(config.ACCIDENTS_DIR.glob(config.ACCIDENTS_FILE_PATTERN),
                         key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise FileNotFoundError(
            f"No BlueKanGo export found matching {config.ACCIDENTS_FILE_PATTERN!r} "
            f"in {config.ACCIDENTS_DIR}/"
        )
    return candidates[0]


def read_working_hours(path: Path = None, target_year: int = None) -> pd.DataFrame:
    """Reads the working hours file (wide format: 1 row per BU incl. a
    "Total" row to skip, 1 column per month) and returns it in long format
    [BU, Year, Month, ID="Saf.4.2", Value]. Months with no value yet (not
    declared) are left out rather than defaulted to 0 or NaN-filled — see
    module docstring. config.WORKING_HOURS_FILE is itself a single-year file
    (e.g. "Working Hours 2026.xlsx", repointed to a new file every January —
    see config.py) — `target_year` just tags the rows it produces, it does
    NOT filter anything; defaults to the current year."""
    target_year = target_year or date.today().year
    path = path or config.WORKING_HOURS_FILE
    if not path.exists():
        raise FileNotFoundError(f"Working hours file not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header = [ws.cell(row=1, column=c).value for c in range(2, ws.max_column + 1)]
    records = []
    for r in range(2, ws.max_row + 1):
        bu = ws.cell(row=r, column=1).value
        if not bu or bu == "Total":
            continue
        for i, month in enumerate(header):
            if month not in MONTHS:
                continue
            value = ws.cell(row=r, column=2 + i).value
            if value is None:
                continue
            records.append({"BU": bu, "Year": target_year, "Month": month, "ID": "Saf.4.2",
                             "Value": float(value), "Comment": None, "Data quality note": None})
    return pd.DataFrame(records)


def _read_accident_events(path: Path = None) -> pd.DataFrame:
    """Reads every individual accident/incident from the latest BlueKanGo
    export, ANY year (the year filter, when needed, is applied by the
    caller — Saf.1 deliberately looks further back than TARGET_YEAR, see
    compute_days_without_accident). Bioline employees only (agency/interim
    workers excluded) and BU already mapped to our internal codes. Returns
    one row per accident: [BU, Date, Lost time (bool), Days lost]."""
    path = path or _find_latest_accidents_file()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    records = []
    unmapped = set()
    for r in range(3, ws.max_row + 1):  # row 1 = title, row 2 = header
        bu_label = ws.cell(row=r, column=2).value
        accident_date = ws.cell(row=r, column=3).value
        interim = ws.cell(row=r, column=4).value
        work_stoppage = ws.cell(row=r, column=5).value
        days_lost = ws.cell(row=r, column=6).value
        if bu_label is None or accident_date is None:
            continue
        if interim == "Oui":
            continue  # not a Bioline employee -> excluded from every KPI definition here
        bu = BU_NAME_MAP.get(bu_label)
        if bu is None:
            unmapped.add(bu_label)
            continue
        records.append({"BU": bu, "Date": accident_date, "LostTime": work_stoppage == "OUI",
                         "DaysLost": days_lost or 0})

    if unmapped:
        print(f"WARNING: unmapped Business Unit label(s) in the BlueKanGo export "
              f"(rows skipped, add them to BU_NAME_MAP): {sorted(unmapped)}")

    df = pd.DataFrame(records, columns=["BU", "Date", "LostTime", "DaysLost"])
    if not df.empty:
        df["Date"] = pd.to_datetime(df["Date"])
    return df


def read_accidents(target_year: int = None, path: Path = None) -> pd.DataFrame:
    """Aggregates accident events (see _read_accident_events) into Saf.2 /
    Saf.3 / Saf.5, per BU and month, restricted to TARGET_YEAR (these three
    are monthly aggregates, unlike Saf.1). Returns one row per
    (BU, Year, Month, ID) in the same long format as the rest of the
    pipeline."""
    target_year = target_year or date.today().year
    events = _read_accident_events(path)
    if events.empty:
        return pd.DataFrame(columns=["BU", "Year", "Month", "ID", "Value", "Comment", "Data quality note"])

    events = events[events["Date"].dt.year == target_year]
    counts = {}  # (bu, month) -> {"Saf.2": int, "Saf.3": int, "Saf.5": float}
    for row in events.itertuples(index=False):
        month = MONTHS[row.Date.month - 1]
        key = (row.BU, month)
        counts.setdefault(key, {"Saf.2": 0, "Saf.3": 0, "Saf.5": 0.0})
        if row.LostTime:
            counts[key]["Saf.3"] += 1
            counts[key]["Saf.5"] += row.DaysLost
        else:
            counts[key]["Saf.2"] += 1

    records = []
    for (bu, month), vals in counts.items():
        for kpi_id in ("Saf.2", "Saf.3", "Saf.5"):
            records.append({"BU": bu, "Year": target_year, "Month": month, "ID": kpi_id, "Value": vals[kpi_id],
                             "Comment": None, "Data quality note": None})
    return pd.DataFrame(records)


def compute_days_without_accident(events: pd.DataFrame, working_hours_df: pd.DataFrame,
                                   target_year: int = None) -> pd.DataFrame:
    """Saf.1 "Number of days without an accident": per Aurélie (22/09/2026),
    a counter since the last accident, updated at the end of each month if
    no accident occurred that month. Computed here as, for each (BU, Year,
    Month) covered by the working hours file (our proxy for "active
    reporting month"): the number of days between the last day of that month
    and the most recent qualifying accident for that BU on or before that
    date — ANY accident (lost-time or not — Aurélie said "no accident", not
    specifically "no lost-time accident"; flag this if it should be
    LTI-only), from ANY year (an accident from a prior year still counts,
    unlike Saf.2/Saf.3/Saf.5 which are restricted to TARGET_YEAR — note the
    Saf.1 row itself is still tagged with the reporting Year of its (BU,
    Month), taken from working_hours_df, NOT the year of the accident it
    counts from). Left out entirely (no row) for a BU with no accident on
    record at all before that month-end — we do not invent a start date to
    count from. `target_year` is only a fallback default for the (BU, Month)
    pairs' Year, used if working_hours_df has no "Year" column (older
    callers) — the normal path reads Year directly off working_hours_df."""
    target_year = target_year or date.today().year
    cols = ["BU", "Month"] + (["Year"] if "Year" in working_hours_df.columns else [])
    active = working_hours_df[cols].drop_duplicates()
    records = []
    for row in active.itertuples(index=False):
        bu, month = row.BU, row.Month
        year = row.Year if "Year" in cols else target_year
        month_num = MONTHS.index(month) + 1
        last_day = date(year, month_num, calendar.monthrange(year, month_num)[1])
        prior = events[(events["BU"] == bu) & (events["Date"].dt.date <= last_day)] if not events.empty \
            else events
        if prior.empty:
            continue
        days = (last_day - prior["Date"].max().date()).days
        records.append({"BU": bu, "Year": year, "Month": month, "ID": "Saf.1", "Value": days,
                         "Comment": None, "Data quality note": None})
    return pd.DataFrame(records)


def fill_zero_accident_months(accidents_df: pd.DataFrame, working_hours_df: pd.DataFrame) -> pd.DataFrame:
    """A BU/year/month with zero qualifying accidents never appears in
    read_accidents' output (nothing to aggregate) — but "no accident" is a
    real value (0), not a missing one. This fills an explicit 0 for
    Saf.2/Saf.3/Saf.5 for every (BU, Year, Month) that IS covered by the
    working hours file and that isn't already present in accidents_df.
    (Saf.1 is handled separately by compute_days_without_accident — it is
    never 0 by default, it needs an actual prior accident date.)"""
    active_months = set(zip(working_hours_df["BU"], working_hours_df["Year"], working_hours_df["Month"]))
    present = set(zip(accidents_df["BU"], accidents_df["Year"], accidents_df["Month"], accidents_df["ID"]))
    filler = []
    for bu, year, month in active_months:
        for kpi_id in ("Saf.2", "Saf.3", "Saf.5"):
            if (bu, year, month, kpi_id) not in present:
                filler.append({"BU": bu, "Year": year, "Month": month, "ID": kpi_id, "Value": 0,
                                "Comment": None, "Data quality note": None})
    if not filler:
        return accidents_df
    return pd.concat([accidents_df, pd.DataFrame(filler)], ignore_index=True)


def main():
    """Entry point: reads the working hours file and the latest BlueKanGo
    export, aggregates them into Saf.1 / Saf.2 / Saf.3 / Saf.4.2 / Saf.5
    rows, and upserts them into input_data/Raw_data_CSR.xlsx (same
    identify-by-(BU,Year,Month,ID), overwrite-existing-rows pattern as
    integrate_data_entry.py — reuses that module's load_consolidated /
    export_consolidated directly rather than duplicating the Excel I/O)."""
    target_year = date.today().year

    working_hours_df = read_working_hours(target_year=target_year)
    print(f"Working hours: {len(working_hours_df)} (BU, Month) values read from "
          f"{config.WORKING_HOURS_FILE.name} (tagged year {target_year})")

    accidents_path = _find_latest_accidents_file()
    events = _read_accident_events(accidents_path)
    accidents_df = read_accidents(target_year=target_year, path=accidents_path)
    accidents_df = fill_zero_accident_months(accidents_df, working_hours_df)
    print(f"Accidents: {accidents_path.name} -> "
          f"{accidents_df[accidents_df.ID == 'Saf.3']['Value'].sum():.0f} LTI, "
          f"{accidents_df[accidents_df.ID == 'Saf.2']['Value'].sum():.0f} NLTI, "
          f"{accidents_df[accidents_df.ID == 'Saf.5']['Value'].sum():.0f} day(s) lost "
          f"(target year {target_year}, agency/interim workers excluded).")

    saf1_df = compute_days_without_accident(events, working_hours_df, target_year=target_year)
    print(f"Days without accident: computed for {len(saf1_df)} (BU, Month) pair(s) "
          f"({len(working_hours_df[['BU', 'Month']].drop_duplicates()) - len(saf1_df)} "
          f"skipped for lack of any accident on record).")

    new_df = pd.concat([working_hours_df, accidents_df, saf1_df], ignore_index=True)

    combined = {(r["BU"], r["Year"], r["Month"], r["ID"]): r for r in load_consolidated().to_dict("records")}
    for r in new_df.to_dict("records"):
        combined[(r["BU"], r["Year"], r["Month"], r["ID"])] = r

    merged = pd.DataFrame(combined.values())
    merged["Month"] = pd.Categorical(merged["Month"], categories=MONTHS, ordered=True)
    merged = merged[["BU", "Year", "Month", "ID", "Value", "Comment", "Data quality note"]]
    merged = merged.sort_values(["BU", "Year", "Month", "ID"]).reset_index(drop=True)
    export_consolidated(merged)
    print(f"\n{config.INPUT_DIR / 'Raw_data_CSR.xlsx'} updated with Safety data "
          f"({len(new_df)} row(s) upserted, year {target_year}).")


if __name__ == "__main__":
    main()
