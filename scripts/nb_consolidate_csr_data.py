# Fabric notebook: nb_consolidate_csr_data
# Workspace: BM_F_D - SAP-B1 / Lakehouse: LH_CSR_Reporting
#
# Portage dans Fabric de la consolidation faite jusqu'ici en local par
# integrate_data_entry.py + integrate_safety_data.py + csr_calc_engine.py
# (+ la partie "détection" de consolidation_report.py, sans l'envoi du mail).
# Même logique, mêmes formules, copiées telles quelles depuis ces scripts —
# seules les entrées/sorties changent :
#   - Lecture : les fichiers SharePoint via le raccourci OneLake
#     Files/sp_csr_general (= Shared Documents/General du site
#     https://biolineagrosciencesgroup.sharepoint.com/sites/CSRreferents),
#     donc toujours la dernière version enregistrée, sans aucune copie.
#   - Raw_data_CSR.xlsx -> table Delta CSR_raw_data (même clé
#     BU/Year/Month/ID, même upsert). Amorcée une seule fois depuis le
#     Raw_data_CSR.xlsx local, pour repartir exactement du même état que le
#     pipeline local (historique 2026 issu des Archives compris).
#   - Consolidated_results_CSR.xlsx -> tables Delta CSR_indicators_report et
#     CSR_completion_tracking.
#   - Rapport d'anomalies -> table CSR_anomalies (append, une colonne
#     Run_date — garde l'historique comme les fichiers horodatés en local).
#     L'email Outlook ne peut pas tourner dans Fabric (pas d'Outlook
#     desktop) : à envoyer depuis le Data Pipeline (activité Office 365
#     Outlook) ou via Activator, à partir de cette table.
#
# Ce notebook ne calcule toujours que des SOMMES, jamais de ratio (décision
# du 23/09/2026, voir csr_calc_engine.py) — les 7 ratios + Env.1 restent
# faits ensuite par nb_consolidate_sap_and_csr_data, à enchaîner juste après
# celui-ci dans le même pipeline.
#
# Double run (local + Fabric) : avec OUTPUT_SUFFIX = "_nb", les tables de
# résultats sont écrites À CÔTÉ de celles produites aujourd'hui par le
# Dataflow (CSR_indicators_report_nb, ...), jamais par-dessus. La cellule 9
# compare au Consolidated_results_CSR.xlsx local. Une fois identique :
# OUTPUT_SUFFIX = "", retirer ces 2 requêtes du Dataflow.
#
# Facteurs de conversion et d'émission : lus dans input_data/CSR_parameters.xlsx
# (un facteur par BU et par année), LE MÊME fichier que le pipeline local —
# rien à recopier ici quand un facteur change. Recopiés en table
# CSR_parameters pour Power BI.
#
# Tant que le pipeline local tourne aussi, les constantes restantes de la
# cellule 2 (BU_COUNTRY, BU_NAME_MAP, seuil) et les fonctions des cellules 3
# à 6 doivent rester synchronisées avec scripts/config.py et les scripts
# d'origine.
#
# Lakehouse par défaut à attacher : LH_CSR_Reporting.

# --------------------------------------------------------------------------
# Cell 0 — dépendance manquante du runtime Fabric (pour Saf.4.1), à mettre
# seule dans sa cellule, sans le "#" :
# --------------------------------------------------------------------------
# %pip install holidays
#
# En exécution planifiée (Data Pipeline), %pip dans le notebook est bloqué
# par défaut — préférer alors un Environment Fabric avec "holidays" en
# bibliothèque publique, attaché à ce notebook.

# --------------------------------------------------------------------------
# Cell 1 — paramètres (à marquer "cellule de paramètres" pour que le
# pipeline puisse les surcharger)
# --------------------------------------------------------------------------
SOURCE_ROOT = "/lakehouse/default/Files/sp_csr_general"
TARGET_YEAR = None      # None = année en cours. Ex. 2026 en janvier 2027 pour finir décembre.
CURRENT_MONTH = None    # None = mois en cours (ou décembre si TARGET_YEAR est une année passée)
OUTPUT_SUFFIX = "_nb"   # "" une fois la bascule faite (voir en-tête)
RESEED_RAW = False      # True = repartir du Raw_data_CSR.xlsx local (écrase CSR_raw_data)

# --------------------------------------------------------------------------
# Cell 2 — imports, chemins, constantes (copiées de scripts/config.py)
# --------------------------------------------------------------------------
import calendar
import glob
import os
from datetime import date

import holidays as holidays_lib
import numpy as np
import openpyxl
import pandas as pd
from pyspark.sql.types import DoubleType, IntegerType, StringType, StructField, StructType

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
MONTH_ORDER = MONTHS

TARGET_YEAR = int(TARGET_YEAR or date.today().year)
if not CURRENT_MONTH:
    CURRENT_MONTH = MONTHS[date.today().month - 1] if TARGET_YEAR == date.today().year else "December"

MONTHLY_REPORTING_DIR = f"{SOURCE_ROOT}/Monthly reporting"
DATA_ENTRY_DIR = MONTHLY_REPORTING_DIR
ACCIDENTS_DIR = f"{MONTHLY_REPORTING_DIR}/extract_bluekango"
ACCIDENTS_FILE_PATTERN = "Accidents_du_travail_Bioline_*.xlsx"
# En local le nom est figé ("Working Hours 2026.xlsx" dans config.py) ; ici
# il suit TARGET_YEAR, donc rien à modifier en janvier tant que le nommage
# reste "Working Hours <année>.xlsx".
WORKING_HOURS_FILE = f"{SOURCE_ROOT}/Working Hours {TARGET_YEAR}.xlsx"
PROJECT_DIR = f"{SOURCE_ROOT}/Reporting_Automation/monthly_reporting"
REFERENCE_FILE = f"{PROJECT_DIR}/input_data/Indicator_reference_CSR.xlsx"
SEED_RAW_FILE = f"{PROJECT_DIR}/input_data/Raw_data_CSR.xlsx"
LOCAL_RESULTS_FILE = f"{PROJECT_DIR}/output_data/Consolidated_results_CSR.xlsx"

RAW_TABLE = "CSR_raw_data"  # pas de suffixe : aucune autre source ne produit cette table
RESULTS_TABLE = f"CSR_indicators_report{OUTPUT_SUFFIX}"
COMPLETION_TABLE = f"CSR_completion_tracking{OUTPUT_SUFFIX}"
ANOMALIES_TABLE = f"CSR_anomalies{OUTPUT_SUFFIX}"

BU_LIST = ["Viridaxis", "BAF", "BFR", "BIB", "BUK", "BUS"]

BU_COUNTRY = {"Viridaxis": "BE", "BAF": "KE", "BFR": "FR", "BIB": "ES", "BUK": "GB", "BUS": "US"}

# Facteurs de conversion et d'émission : PAS ici, mais dans
# input_data/CSR_parameters.xlsx (Parameter, BU, Year, Value, Unit, Source),
# le même fichier que lit le pipeline local — voir la cellule 5 et
# scripts/csr_parameters.py.
PARAMETERS_FILE = f"{PROJECT_DIR}/input_data/CSR_parameters.xlsx"
PARAMETERS_TABLE = f"CSR_parameters{OUTPUT_SUFFIX}"

VARIATION_THRESHOLD = 0.20

# Libellé "Business Unit" de BlueKanGo -> code BU interne (integrate_safety_data.py).
BU_NAME_MAP = {
    "Bioline Africa": "BAF",
    "Bioline US": "BUS",
    "Bioline UK": "BUK",
    "Bioline Iberia": "BIB",
    "Bioline France": "BFR",
    "Bioline Viridaxis": "Viridaxis",  # non vérifié — aucun exemple vu dans l'export
}

RAW_COLS = ["BU", "Year", "Month", "ID", "Value", "Comment", "Data quality note"]

# --------------------------------------------------------------------------
# Cell 3 — lecture des fichiers de saisie des 6 BU (integrate_data_entry.py)
# --------------------------------------------------------------------------
import re

FRENCH_DECIMAL_RE = re.compile(r"^\s*-?\d+,\d+\s*$")

# Colonnes du fichier de saisie : ID, puis les 6 colonnes de contexte de
# generate_data_entry_file.DISPLAY_COLS, puis "Value of the month" et
# "Comment" -> colonnes 8 et 9.
DATA_ENTRY_VALUE_COL = 6 + 2
DATA_ENTRY_COMMENT_COL = DATA_ENTRY_VALUE_COL + 1


def parse_raw_value(v):
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None, None
        if FRENCH_DECIMAL_RE.match(s):
            return float(s.replace(",", ".")), "Corrected: decimal comma (“" + v + "”) -> dot"
        try:
            return float(s), "Corrected: numeric text -> number"
        except ValueError:
            return None, f"NOT PARSABLE, needs a manual fix in the data entry file: “{v}”"
    return None, f"Unexpected type ({type(v).__name__}): {v!r}"


def read_data_entry_file(bu: str, year: int) -> pd.DataFrame:
    path = f"{DATA_ENTRY_DIR}/{bu}_data_entry_CSR_{year}.xlsm"
    if not os.path.exists(path):
        raise FileNotFoundError(f"Data entry file not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    records = []
    for month in MONTHS:
        if month not in wb.sheetnames:
            continue
        ws = wb[month]
        for r in range(2, ws.max_row + 1):
            idv = ws.cell(row=r, column=1).value  # "ID" column
            if not idv:
                continue
            raw_value = ws.cell(row=r, column=DATA_ENTRY_VALUE_COL).value
            comment = ws.cell(row=r, column=DATA_ENTRY_COMMENT_COL).value
            value, note = parse_raw_value(raw_value)
            records.append({"BU": bu, "Year": year, "Month": month, "ID": idv, "Value": value,
                            "Comment": comment, "Data quality note": note})
    return pd.DataFrame(records, columns=RAW_COLS)

# --------------------------------------------------------------------------
# Cell 4 — Safety : Working Hours + BlueKanGo (integrate_safety_data.py)
# --------------------------------------------------------------------------


def _find_latest_accidents_file() -> str:
    """Le dernier export d'après l'horodatage DANS LE NOM
    (..._20260921-165507.xlsx, trié lexicographiquement) — plus fiable à
    travers un raccourci OneLake que la date de modification utilisée en
    local."""
    candidates = sorted(glob.glob(f"{ACCIDENTS_DIR}/{ACCIDENTS_FILE_PATTERN}"), key=os.path.basename)
    if not candidates:
        raise FileNotFoundError(f"No BlueKanGo export found matching {ACCIDENTS_FILE_PATTERN!r} in {ACCIDENTS_DIR}/")
    return candidates[-1]


def read_working_hours(path: str, target_year: int) -> pd.DataFrame:
    if not os.path.exists(path):
        raise FileNotFoundError(f"Working hours file not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header = [ws.cell(row=1, column=c).value for c in range(2, ws.max_column + 1)]
    records = []
    for r in range(2, ws.max_row + 1):
        bu = ws.cell(row=r, column=1).value
        if not bu or bu == "Total":
            continue
        for i, month in enumerate(header):
            if month not in MONTHS:
                continue
            value = ws.cell(row=r, column=2 + i).value
            if value is None:
                continue
            records.append({"BU": bu, "Year": target_year, "Month": month, "ID": "Saf.4.2",
                            "Value": float(value), "Comment": None, "Data quality note": None})
    return pd.DataFrame(records, columns=RAW_COLS)


def _read_accident_events(path: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    records = []
    unmapped = set()
    for r in range(3, ws.max_row + 1):  # row 1 = title, row 2 = header
        bu_label = ws.cell(row=r, column=2).value
        accident_date = ws.cell(row=r, column=3).value
        interim = ws.cell(row=r, column=4).value
        work_stoppage = ws.cell(row=r, column=5).value
        days_lost = ws.cell(row=r, column=6).value
        if bu_label is None or accident_date is None:
            continue
        if interim == "Oui":
            continue  # not a Bioline employee -> excluded from every KPI definition here
        bu = BU_NAME_MAP.get(bu_label)
        if bu is None:
            unmapped.add(bu_label)
            continue
        records.append({"BU": bu, "Date": accident_date, "LostTime": work_stoppage == "OUI",
                        "DaysLost": days_lost or 0})

    if unmapped:
        print(f"WARNING: unmapped Business Unit label(s) in the BlueKanGo export "
              f"(rows skipped, add them to BU_NAME_MAP): {sorted(unmapped)}")

    df = pd.DataFrame(records, columns=["BU", "Date", "LostTime", "DaysLost"])
    if not df.empty:
        df["Date"] = pd.to_datetime(df["Date"])
    return df


def read_accidents(events: pd.DataFrame, target_year: int) -> pd.DataFrame:
    if events.empty:
        return pd.DataFrame(columns=RAW_COLS)

    events = events[events["Date"].dt.year == target_year]
    counts = {}  # (bu, month) -> {"Saf.2": int, "Saf.3": int, "Saf.5": float}
    for row in events.itertuples(index=False):
        month = MONTHS[row.Date.month - 1]
        key = (row.BU, month)
        counts.setdefault(key, {"Saf.2": 0, "Saf.3": 0, "Saf.5": 0.0})
        if row.LostTime:
            counts[key]["Saf.3"] += 1
            counts[key]["Saf.5"] += row.DaysLost
        else:
            counts[key]["Saf.2"] += 1

    records = []
    for (bu, month), vals in counts.items():
        for kpi_id in ("Saf.2", "Saf.3", "Saf.5"):
            records.append({"BU": bu, "Year": target_year, "Month": month, "ID": kpi_id, "Value": vals[kpi_id],
                            "Comment": None, "Data quality note": None})
    return pd.DataFrame(records, columns=RAW_COLS)


def compute_days_without_accident(events: pd.DataFrame, working_hours_df: pd.DataFrame) -> pd.DataFrame:
    """Saf.1 : jours entre la fin du mois et le dernier accident (tout type,
    toute année) de la BU — voir integrate_safety_data.py pour le détail."""
    active = working_hours_df[["BU", "Month", "Year"]].drop_duplicates()
    records = []
    for row in active.itertuples(index=False):
        bu, month, year = row.BU, row.Month, int(row.Year)
        month_num = MONTHS.index(month) + 1
        last_day = date(year, month_num, calendar.monthrange(year, month_num)[1])
        prior = events[(events["BU"] == bu) & (events["Date"].dt.date <= last_day)] if not events.empty \
            else events
        if prior.empty:
            continue
        days = (last_day - prior["Date"].max().date()).days
        records.append({"BU": bu, "Year": year, "Month": month, "ID": "Saf.1", "Value": days,
                        "Comment": None, "Data quality note": None})
    return pd.DataFrame(records, columns=RAW_COLS)


def fill_zero_accident_months(accidents_df: pd.DataFrame, working_hours_df: pd.DataFrame) -> pd.DataFrame:
    active_months = set(zip(working_hours_df["BU"], working_hours_df["Year"], working_hours_df["Month"]))
    present = set(zip(accidents_df["BU"], accidents_df["Year"], accidents_df["Month"], accidents_df["ID"]))
    filler = []
    for bu, year, month in active_months:
        for kpi_id in ("Saf.2", "Saf.3", "Saf.5"):
            if (bu, year, month, kpi_id) not in present:
                filler.append({"BU": bu, "Year": year, "Month": month, "ID": kpi_id, "Value": 0,
                               "Comment": None, "Data quality note": None})
    if not filler:
        return accidents_df
    return pd.concat([accidents_df, pd.DataFrame(filler, columns=RAW_COLS)], ignore_index=True)

# --------------------------------------------------------------------------
# Cell 5 — moteur de calcul (csr_calc_engine.py) — SOMMES uniquement
# --------------------------------------------------------------------------


def z(x):
    return 0 if x is None or (isinstance(x, float) and pd.isna(x)) else x


FORMULAS = {
    "Ene.9": lambda v: (z(v("Ene.1")) + z(v("Ene.2")) + z(v("Ene.3")) + z(v("Ene.4")) + z(v("Ene.5"))
                        + z(v("Ene.6")) * z(v("Ene.6.1"))
                        + z(v("Ene.7")) * z(v("Ene.7.1"))
                        + z(v("Ene.8")) * z(v("Ene.7.1"))),
    "Ref.1": lambda v: sum(z(v(f"Ref.{i}")) for i in range(2, 10)),
    "Carb.1": lambda v: z(v("Ene.1")) * z(v("Ene.12")),
    "Carb.2": lambda v: z(v("Ene.5")) * z(v("Ene.13")),
    "Carb.3": lambda v: z(v("Ene.7")) * z(v("Ene.14")),
    "Carb.4": lambda v: z(v("Ene.6")) * z(v("Ene.15")),
    "Carb.5": lambda v: sum(z(v(f"Carb.{i}")) for i in range(1, 5)),
}


def working_days_in_month(country_code: str, year: int, month: int) -> int:
    country_holidays = holidays_lib.country_holidays(country_code, years=year)
    _, days_in_month = calendar.monthrange(year, month)
    return sum(
        1 for day in range(1, days_in_month + 1)
        if date(year, month, day).weekday() < 5 and date(year, month, day) not in country_holidays
    )


# --- Paramètres (scripts/csr_parameters.py) : un facteur par (BU, Year),
# repli sur l'année antérieure la plus récente, jamais sur une année
# postérieure, signalé dans les anomalies.
PARAMETER_COLS = ["Parameter", "BU", "Year", "Value", "Unit", "Source"]
PARAMETER_FOR_ID = {
    "Ene.6.1": "LPG_CONVERSION_FACTOR",
    "Ene.7.1": "FUEL_CONVERSION_FACTOR",
    "Ene.12": "ELECTRICITY_EMISSION_FACTOR",
    "Ene.13": "NATURAL_GAS_EMISSION_FACTOR",
    "Ene.14": "FUEL_EMISSION_FACTOR",
    "Ene.15": "LPG_EMISSION_FACTOR",
}


class MissingParameterError(LookupError):
    pass


def validate_parameters(parameters: pd.DataFrame) -> None:
    missing_cols = [c for c in ["Parameter", "BU", "Year", "Value"] if c not in parameters.columns]
    if missing_cols:
        raise ValueError(f"CSR_parameters.xlsx: missing column(s) {missing_cols}")
    empty = parameters[parameters["Value"].isna()]
    if len(empty):
        raise ValueError(f"CSR_parameters.xlsx: empty Value on row(s):\n{empty.to_string(index=False)}")
    dupes = parameters[parameters.duplicated(["Parameter", "BU", "Year"], keep=False)]
    if len(dupes):
        raise ValueError(f"CSR_parameters.xlsx: more than one row for the same (Parameter, BU, Year):\n"
                         f"{dupes.to_string(index=False)}")


def build_parameter_index(parameters) -> dict:
    index = {}
    if parameters is None:
        return index
    for name, bu, year, value in parameters[["Parameter", "BU", "Year", "Value"]].itertuples(index=False):
        index.setdefault((name, bu), []).append((int(year), float(value)))
    for rows in index.values():
        rows.sort()
    return index


def resolve_parameter(index: dict, name: str, bu: str, year: int):
    candidates = [row for row in index.get((name, bu), []) if row[0] <= year]
    if not candidates:
        raise MissingParameterError(f"{name} for {bu}: no value for {year} or any earlier year")
    year_used, value = candidates[-1]
    return value, year_used


def find_parameter_issues(parameters, raw: pd.DataFrame) -> pd.DataFrame:
    index = build_parameter_index(parameters)
    rows = []
    for bu, year in sorted(set(zip(raw["BU"], raw["Year"].astype(int)))):
        for name in sorted(set(PARAMETER_FOR_ID.values())):
            try:
                _, year_used = resolve_parameter(index, name, bu, year)
            except MissingParameterError:
                rows.append({"BU": bu, "Year": year, "Parameter": name,
                             "Issue": "Missing: dependent indicators not computed"})
                continue
            if year_used != year:
                rows.append({"BU": bu, "Year": year, "Parameter": name,
                             "Issue": f"No {year} value, {year_used} value used instead"})
    return pd.DataFrame(rows, columns=["BU", "Year", "Parameter", "Issue"])


def _parameter(name):
    def resolve(bu, year, month, parameters):
        return resolve_parameter(parameters, name, bu, year)[0]
    return resolve


CONTEXTUAL_VALUES = {
    **{kpi_id: _parameter(name) for kpi_id, name in PARAMETER_FOR_ID.items()},
    "Saf.4.1": lambda bu, year, month, parameters: working_days_in_month(
        BU_COUNTRY[bu], year, MONTH_ORDER.index(month) + 1),
}


def compute_bu_month(raw_values: dict) -> dict:
    values = dict(raw_values)
    remaining = dict(FORMULAS)
    progress = True
    while remaining and progress:
        progress = False
        for kpi_id, formula in list(remaining.items()):
            try:
                values[kpi_id] = formula(lambda i: values[i])
                del remaining[kpi_id]
                progress = True
            except (KeyError, ZeroDivisionError, TypeError):
                continue
    return values


def run(reference: pd.DataFrame, raw: pd.DataFrame, parameters: pd.DataFrame = None) -> pd.DataFrame:
    meta = reference.set_index("ID")[["Topic", "KPI", "Unit", "Kind"]]
    parameter_index = build_parameter_index(parameters)
    results = []
    for (bu, year, month), grp in raw.groupby(["BU", "Year", "Month"], sort=False):
        year = int(year)
        raw_values = dict(zip(grp["ID"], grp["Value"]))
        for kpi_id, contextual_fn in CONTEXTUAL_VALUES.items():
            try:
                raw_values[kpi_id] = contextual_fn(bu, year, month, parameter_index)
            except MissingParameterError:
                continue  # facteur absent -> indicateurs dépendants non calculés, signalé dans CSR_anomalies
        computed = compute_bu_month(raw_values)
        for kpi_id, value in computed.items():
            if kpi_id not in meta.index:
                continue
            row = meta.loc[kpi_id]
            results.append({
                "BU": bu, "Year": year, "Month": month, "ID": kpi_id,
                "Topic": row["Topic"], "KPI": row["KPI"], "Unit": row["Unit"],
                "Kind": row["Kind"], "Value": value,
            })
    out = pd.DataFrame(results)
    out["Month"] = pd.Categorical(out["Month"], categories=MONTH_ORDER, ordered=True)
    return out.sort_values(["BU", "Year", "Month", "ID"]).reset_index(drop=True)


def build_completion_table(raw: pd.DataFrame, reference: pd.DataFrame, current_year: int) -> pd.DataFrame:
    raw = raw[raw.Year == current_year]
    referent_ids = set(reference[(reference.Kind == "input") & (reference.Responsible == "CSR referent")]["ID"])
    total = len(referent_ids)

    filled = raw[raw.ID.isin(referent_ids) & raw["Value"].notna()]
    filled_counts = filled.groupby(["BU", "Month"]).size()

    rows = []
    for bu in sorted(raw["BU"].unique()):
        for month in MONTH_ORDER:
            count = filled_counts.get((bu, month), 0)
            rows.append({
                "BU": bu,
                "Month": month,
                "completion_percent": round(100 * count / total) if total else 0,
            })
    return pd.DataFrame(rows)

# --------------------------------------------------------------------------
# Cell 6 — détection des anomalies (consolidation_report.py, sans l'email)
# --------------------------------------------------------------------------


def find_data_quality_notes(raw: pd.DataFrame) -> pd.DataFrame:
    return raw[raw["Data quality note"].notna()][["BU", "Year", "Month", "ID", "Value", "Data quality note"]]


def find_missing_values(raw: pd.DataFrame, input_ids: set, current_year: int, current_month: str) -> pd.DataFrame:
    raw = raw[raw.Year == current_year]
    due_months = MONTH_ORDER[:MONTH_ORDER.index(current_month) + 1]
    bus = sorted(raw["BU"].unique())
    if not bus or not input_ids:
        return pd.DataFrame(columns=["BU", "Month", "ID"])

    expected = pd.MultiIndex.from_product([bus, due_months, sorted(input_ids)],
                                          names=["BU", "Month", "ID"]).to_frame(index=False)
    have_value = raw[raw.Month.isin(due_months) & raw.ID.isin(input_ids) & raw["Value"].notna()]
    merged = expected.merge(have_value[["BU", "Month", "ID"]], on=["BU", "Month", "ID"],
                            how="left", indicator=True)
    missing = merged[merged["_merge"] == "left_only"][["BU", "Month", "ID"]].copy()
    missing["Month"] = pd.Categorical(missing["Month"], categories=MONTH_ORDER, ordered=True)
    return missing.sort_values(["BU", "Month", "ID"]).reset_index(drop=True)


def find_large_variations(consolidated: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (bu, year, kpi_id), grp in consolidated.groupby(["BU", "Year", "ID"], sort=False):
        grp = grp.sort_values("Month")
        prev_value, prev_month = None, None
        for _, r in grp.iterrows():
            value = r["Value"]
            if pd.notna(value) and prev_value is not None and prev_value != 0:
                variation = abs(value - prev_value) / abs(prev_value)
                if variation > VARIATION_THRESHOLD:
                    rows.append({
                        "BU": bu, "Year": year, "ID": kpi_id,
                        "From": prev_month, "From value": prev_value,
                        "To": r["Month"], "To value": value,
                        "Variation": f"{variation:.0%}",
                    })
            if pd.notna(value):
                prev_value, prev_month = value, r["Month"]
    return pd.DataFrame(rows, columns=["BU", "Year", "ID", "From", "From value", "To", "To value", "Variation"])


def build_anomalies_table(raw: pd.DataFrame, consolidated: pd.DataFrame, reference: pd.DataFrame,
                          current_year: int, current_month: str, parameters: pd.DataFrame) -> pd.DataFrame:
    """Les 4 contrôles du rapport local, empilés dans UNE table longue
    (Check, BU, Year, Month, ID, Value, Detail) — c'est elle que le pipeline
    lira pour composer le mail (compte par BU = un simple groupby)."""
    input_ids = set(reference[reference.Kind == "input"]["ID"])
    quality = find_data_quality_notes(raw)
    missing = find_missing_values(raw, input_ids, current_year, current_month)
    variations = find_large_variations(consolidated)
    param_issues = find_parameter_issues(parameters, raw)

    parts = [
        pd.DataFrame({"Check": "Data quality", "BU": quality["BU"], "Year": quality["Year"],
                      "Month": quality["Month"], "ID": quality["ID"], "Value": quality["Value"],
                      "Detail": quality["Data quality note"]}),
        pd.DataFrame({"Check": "Missing value", "BU": missing["BU"], "Year": current_year,
                      "Month": missing["Month"].astype(str), "ID": missing["ID"], "Value": np.nan,
                      "Detail": "No value entered"}),
        pd.DataFrame({"Check": "Large variation", "BU": variations["BU"], "Year": variations["Year"],
                      "Month": variations["To"].astype(str), "ID": variations["ID"],
                      "Value": variations["To value"],
                      "Detail": (variations["From"].astype(str) + ": " + variations["From value"].astype(str)
                                 + " -> " + variations["To"].astype(str) + ": "
                                 + variations["To value"].astype(str) + " (" + variations["Variation"] + ")")}),
        pd.DataFrame({"Check": "Parameter", "BU": param_issues["BU"], "Year": param_issues["Year"],
                      "Month": None, "ID": param_issues["Parameter"], "Value": np.nan,
                      "Detail": param_issues["Issue"]}),
    ]
    out = pd.concat([p for p in parts if not p.empty], ignore_index=True) if any(not p.empty for p in parts) \
        else pd.DataFrame(columns=["Check", "BU", "Year", "Month", "ID", "Value", "Detail"])
    out["Run_date"] = date.today().isoformat()
    return out

# --------------------------------------------------------------------------
# Cell 7 — lecture/écriture Delta
# --------------------------------------------------------------------------
# Delta refuse les espaces dans les noms de colonnes : "Data quality note"
# devient "Data_quality_note" dans la table, et redevient "Data quality note"
# à la relecture, pour que les fonctions copiées ci-dessus restent
# identiques aux scripts locaux.
RAW_SCHEMA = StructType([
    StructField("BU", StringType()), StructField("Year", IntegerType()),
    StructField("Month", StringType()), StructField("ID", StringType()),
    StructField("Value", DoubleType()), StructField("Comment", StringType()),
    StructField("Data_quality_note", StringType()),
])
RESULTS_SCHEMA = StructType([
    StructField("BU", StringType()), StructField("Year", IntegerType()),
    StructField("Month", StringType()), StructField("ID", StringType()),
    StructField("Topic", StringType()), StructField("KPI", StringType()),
    StructField("Unit", StringType()), StructField("Kind", StringType()),
    StructField("Value", DoubleType()),
])
COMPLETION_SCHEMA = StructType([
    StructField("BU", StringType()), StructField("Year", IntegerType()),
    StructField("Month", StringType()), StructField("completion_percent", IntegerType()),
])
PARAMETERS_SCHEMA = StructType([
    StructField("Parameter", StringType()), StructField("BU", StringType()),
    StructField("Year", IntegerType()), StructField("Value", DoubleType()),
    StructField("Unit", StringType()), StructField("Source", StringType()),
])
ANOMALIES_SCHEMA = StructType([
    StructField("Check", StringType()), StructField("BU", StringType()),
    StructField("Year", IntegerType()), StructField("Month", StringType()),
    StructField("ID", StringType()), StructField("Value", DoubleType()),
    StructField("Detail", StringType()), StructField("Run_date", StringType()),
])

_CASTS = {StringType: str, IntegerType: int, DoubleType: float}


def write_table(df: pd.DataFrame, name: str, schema: StructType, mode: str = "overwrite"):
    """pandas -> Spark avec un schéma explicite : sans lui, l'inférence de
    type échoue sur les colonnes mixtes (Comment : texte, nombre ou vide)."""
    df = df.rename(columns={"Data quality note": "Data_quality_note"})
    rows = []
    for rec in df[[f.name for f in schema.fields]].astype(object).itertuples(index=False):
        rows.append(tuple(
            None if v is None or (isinstance(v, float) and pd.isna(v)) else _CASTS[type(f.dataType)](v)
            for v, f in zip(rec, schema.fields)
        ))
    writer = spark.createDataFrame(rows, schema).write.format("delta").mode(mode)
    if mode == "overwrite":
        writer = writer.option("overwriteSchema", "true")
    writer.saveAsTable(name)
    print(f"{name}: {len(rows)} row(s) written ({mode}).")


def load_raw() -> pd.DataFrame:
    if spark.catalog.tableExists(RAW_TABLE) and not RESEED_RAW:
        raw = spark.read.table(RAW_TABLE).toPandas().rename(columns={"Data_quality_note": "Data quality note"})
        print(f"{RAW_TABLE}: {len(raw)} row(s) read.")
    else:
        raw = pd.read_excel(SEED_RAW_FILE, sheet_name="Raw_data")
        print(f"{RAW_TABLE}: seeded from {SEED_RAW_FILE} ({len(raw)} row(s)).")
    raw = raw[RAW_COLS].copy()
    raw["Year"] = raw["Year"].astype("int64")
    raw["Month"] = raw["Month"].astype(str)
    return raw


def upsert(existing: pd.DataFrame, new: pd.DataFrame) -> pd.DataFrame:
    """Même upsert que integrate_data_entry.py / integrate_safety_data.py :
    clé (BU, Year, Month, ID), la nouvelle ligne remplace l'ancienne, tout le
    reste (dont les autres années) est conservé."""
    combined = {(r["BU"], int(r["Year"]), r["Month"], r["ID"]): r for r in existing.to_dict("records")}
    for r in new.to_dict("records"):
        combined[(r["BU"], int(r["Year"]), r["Month"], r["ID"])] = r
    merged = pd.DataFrame(list(combined.values()), columns=RAW_COLS)
    merged["Year"] = merged["Year"].astype("int64")
    merged["_m"] = pd.Categorical(merged["Month"], categories=MONTHS, ordered=True)
    return merged.sort_values(["BU", "Year", "_m", "ID"]).drop(columns="_m").reset_index(drop=True)

# --------------------------------------------------------------------------
# Cell 8 — exécution : intégration des 6 BU + Safety, calcul, écriture
# --------------------------------------------------------------------------
reference = pd.read_excel(REFERENCE_FILE, sheet_name="Reference")
parameters = pd.read_excel(PARAMETERS_FILE, sheet_name="Parameters")
validate_parameters(parameters)
raw = load_raw()

# 1. Fichiers de saisie des référents (integrate_data_entry.py)
entries = []
for bu in BU_LIST:
    try:
        df = read_data_entry_file(bu, TARGET_YEAR)
    except FileNotFoundError as e:
        print(f"{bu}: {e}")
        continue
    flagged = df[df["Data quality note"].notna()]
    print(f"{bu}: {len(df)} values read from the data entry file"
          + (f", {len(flagged)} quality note(s)" if len(flagged) else ""))
    entries.append(df)
if entries:
    raw = upsert(raw, pd.concat(entries, ignore_index=True))

# 2. Safety (integrate_safety_data.py)
working_hours_df = read_working_hours(WORKING_HOURS_FILE, TARGET_YEAR)
accidents_path = _find_latest_accidents_file()
events = _read_accident_events(accidents_path)
accidents_df = fill_zero_accident_months(read_accidents(events, TARGET_YEAR), working_hours_df)
saf1_df = compute_days_without_accident(events, working_hours_df)
print(f"Working hours: {len(working_hours_df)} value(s) from {os.path.basename(WORKING_HOURS_FILE)}")
print(f"Accidents: {os.path.basename(accidents_path)} -> "
      f"{accidents_df[accidents_df.ID == 'Saf.3']['Value'].sum():.0f} LTI, "
      f"{accidents_df[accidents_df.ID == 'Saf.2']['Value'].sum():.0f} NLTI, "
      f"{accidents_df[accidents_df.ID == 'Saf.5']['Value'].sum():.0f} day(s) lost ({TARGET_YEAR})")
raw = upsert(raw, pd.concat([working_hours_df, accidents_df, saf1_df], ignore_index=True))

write_table(raw, RAW_TABLE, RAW_SCHEMA)

# 3. Calcul (csr_calc_engine.py)
consolidated = run(reference, raw, parameters)
completion = build_completion_table(raw, reference, TARGET_YEAR)
completion.insert(1, "Year", TARGET_YEAR)
anomalies = build_anomalies_table(raw, consolidated, reference, TARGET_YEAR, CURRENT_MONTH, parameters)

print(f"\n{len(consolidated)} rows computed for {raw.BU.nunique()} BUs, {raw.Year.nunique()} year(s) "
      f"({consolidated.Value.notna().sum()} non-empty values).")
print("\nAnomalies per BU:")
print(anomalies.groupby(["BU", "Check"]).size().unstack(fill_value=0).to_string() if len(anomalies)
      else "No anomaly found.")

write_table(consolidated.assign(Month=consolidated["Month"].astype(str)), RESULTS_TABLE, RESULTS_SCHEMA)
write_table(completion, COMPLETION_TABLE, COMPLETION_SCHEMA)
write_table(parameters.reindex(columns=PARAMETER_COLS), PARAMETERS_TABLE, PARAMETERS_SCHEMA)  # pour Power BI
write_table(anomalies, ANOMALIES_TABLE, ANOMALIES_SCHEMA, mode="append")

# --------------------------------------------------------------------------
# Cell 9 — double run : comparaison au Consolidated_results_CSR.xlsx local
# --------------------------------------------------------------------------
# À lancer tant que OUTPUT_SUFFIX = "_nb". Un écart n'est pas forcément un
# bug : si un référent a saisi depuis le dernier run local, Fabric (qui lit
# la version à jour) a une valeur que le fichier local n'a pas encore —
# relancer le pipeline local juste avant pour comparer à état égal.
keys = ["BU", "Year", "Month", "ID"]
local = pd.read_excel(LOCAL_RESULTS_FILE, sheet_name="Results")[keys + ["Value"]]
local["Month"] = local["Month"].astype(str)
fabric = consolidated[keys + ["Value"]].assign(Month=consolidated["Month"].astype(str))

cmp = local.merge(fabric, on=keys, how="outer", suffixes=("_local", "_fabric"), indicator=True)
a, b = cmp["Value_local"], cmp["Value_fabric"]
same = (a.isna() & b.isna()) | ((a - b).abs() <= 1e-9 * np.maximum(1, a.abs()))
diff = cmp[(cmp["_merge"] != "both") | ~same]

print(f"{len(cmp)} (BU, Year, Month, ID) compared, {len(diff)} difference(s).")
if len(diff):
    print(diff.groupby(["_merge", "ID"]).size().to_string())
    print(diff.head(50).to_string(index=False))
