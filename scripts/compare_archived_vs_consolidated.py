"""
Compares indicator values as they ORIGINALLY appeared in the raw
12-tab-per-BU workbooks (config.RAW_DATA_DIR, the "Archives" folder) against
what's now in output_data/Consolidated_results_CSR.xlsx, to make it easy to
see, concretely, how much a pipeline change actually moved the numbers —
e.g. after switching Safety data from hand-typed values to BlueKanGo/Working
Hours (22-23/09/2026), or after any future change to FORMULAS/
CONTEXTUAL_VALUES/EXCLUDED_IDS.

NOT part of the recurring monthly cycle — a diagnostic tool, run by hand
whenever you want to sanity-check a change. Only ever compares
HISTORICAL_YEAR (2026, see extract_reference_and_data.py): the original raw
workbooks never covered any other year, so comparing them against a later
year in Consolidated_results_CSR.xlsx wouldn't mean anything.

Three kinds of difference, and this tool tells them apart instead of just
dumping every mismatch as if they were equally alarming:
  - "expected_absence": an ID that's supposed to be missing from the
    consolidated output by design — derived live from
    extract_reference_and_data.EXCLUDED_IDS (Env.1, Saf.4 — never tracked at
    all anymore) and from extract_reference_and_data.CALC_IDS minus
    csr_calc_engine's FORMULAS/CONTEXTUAL_VALUES (the ratio indicators, no
    longer computed by this engine — see csr_calc_engine.py's docstring).
    Never hardcoded as a fixed list: if someone adds a formula back, or
    excludes a new ID, this tool picks it up automatically on the next run.
  - "unexpected_absence" / "unexpected_new_only": an ID missing (or newly
    appearing) that ISN'T explained by a known exclusion — worth a look.
  - "differs": present on both sides, with a different value — normal and
    expected for Safety (Saf.1/2/3/4.2/5, now sourced from BlueKanGo/Working
    Hours instead of a hand-typed snapshot), worth a second look for
    anything else.

Usage
-----
    uv run scripts/compare_archived_vs_consolidated.py            # every indicator
    uv run scripts/compare_archived_vs_consolidated.py Saf         # only IDs starting with "Saf"
    uv run scripts/compare_archived_vs_consolidated.py Saf 10      # up to 10 sample rows per ID (default 5)

Besides the console summary, every run also writes a timestamped workbook to
output_data/Comparison_archived_vs_consolidated_<date>.xlsx (two sheets:
"Summary", one row per ID; "Details", every row not a plain match or
expected absence) — the console output only ever shows a sample.

Prerequisites: the 6 original raw files still reachable at
config.RAW_DATA_DIR, and output_data/Consolidated_results_CSR.xlsx already
generated (uv run scripts/csr_calc_engine.py).

Dependencies: pandas, openpyxl.
"""
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config
import csr_calc_engine
import extract_reference_and_data as extract

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


def read_original_indicators(bu: str) -> pd.DataFrame:
    """Every indicator row (any ID, any Kind — unlike
    extract_reference_and_data.build_raw_long, which only keeps "input"
    ones) in this BU's original raw workbook: [BU, Month, ID, Original].
    Reading everything, not just what's still tracked today, is the whole
    point here — we want to see what changed, including things that
    disappeared entirely."""
    path = config.RAW_DATA_DIR / config.RAW_FILES[bu]
    wb = openpyxl.load_workbook(path, data_only=True)
    records = []
    for month in MONTHS:
        if month not in wb.sheetnames:
            continue
        ws = wb[month]
        for r in range(2, ws.max_row + 1):
            idv = ws.cell(row=r, column=2).value
            if not idv:
                continue
            records.append({"BU": bu, "Month": month, "ID": idv,
                             "Original": ws.cell(row=r, column=10).value})
    return pd.DataFrame(records)


def load_all_original() -> pd.DataFrame:
    """read_original_indicators for every BU in config.RAW_FILES, combined."""
    return pd.concat([read_original_indicators(bu) for bu in config.RAW_FILES], ignore_index=True)


def load_consolidated(year: int) -> pd.DataFrame:
    """output_data/Consolidated_results_CSR.xlsx's "Results" sheet, restricted
    to `year` (only HISTORICAL_YEAR ever makes sense to compare — see module
    docstring)."""
    path = config.OUTPUT_DIR / "Consolidated_results_CSR.xlsx"
    if not path.exists():
        raise FileNotFoundError(
            f"{path} not found — run `uv run scripts/csr_calc_engine.py` first."
        )
    df = pd.read_excel(path, sheet_name="Results")
    return df[df.Year == year][["BU", "Month", "ID", "Value"]].rename(columns={"Value": "New"})


def known_absent_from_consolidated() -> set:
    """IDs that are SUPPOSED to be missing from Consolidated_results_CSR.xlsx
    by design, derived live from the actual current state of the pipeline
    rather than a fixed list that could silently drift out of sync:
      - extract_reference_and_data.EXCLUDED_IDS: never even reach the
        reference list (Env.1, Saf.4 as of 23/09/2026).
      - extract_reference_and_data.CALC_IDS minus whatever
        csr_calc_engine.py actually produces (its FORMULAS and
        CONTEXTUAL_VALUES combined): indicators still classified
        "calculated" in the reference list, but not computed by this
        engine (the ratio indicators, computed downstream instead)."""
    produced_by_engine = set(csr_calc_engine.FORMULAS) | set(csr_calc_engine.CONTEXTUAL_VALUES)
    return extract.EXCLUDED_IDS | (extract.CALC_IDS - produced_by_engine)


def compare(original: pd.DataFrame, consolidated: pd.DataFrame, known_absent: set = None) -> pd.DataFrame:
    """Merges `original` (BU, Month, ID, Original) and `consolidated` (BU,
    Month, ID, New) on (BU, Month, ID), and classifies every resulting row
    into a Status:
      - "match": same value on both sides (numerically equal, or both
        empty/NaN).
      - "differs": present on both sides, different value.
      - "expected_absence": only in `original`, and ID is in `known_absent`.
      - "unexpected_absence": only in `original`, NOT in `known_absent` —
        worth investigating, this ID should still be produced.
      - "unexpected_new_only": only in `consolidated` — an ID that
        materialized with no trace in the original archive at all."""
    known_absent = known_absent if known_absent is not None else set()
    merged = original.merge(consolidated, on=["BU", "Month", "ID"], how="outer", indicator=True)
    merged["Original_num"] = pd.to_numeric(merged["Original"], errors="coerce")
    merged["New_num"] = pd.to_numeric(merged["New"], errors="coerce")

    def status(row):
        if row["_merge"] == "left_only":
            return "expected_absence" if row["ID"] in known_absent else "unexpected_absence"
        if row["_merge"] == "right_only":
            return "unexpected_new_only"
        both_empty = pd.isna(row["Original_num"]) and pd.isna(row["New_num"])
        same = both_empty or round(row["Original_num"], 6) == round(row["New_num"], 6)
        return "match" if same else "differs"

    merged["Status"] = merged.apply(status, axis=1)
    return merged[["BU", "Month", "ID", "Original", "New", "Status"]]


def summarize(compared: pd.DataFrame) -> pd.DataFrame:
    """One row per ID, counting each Status — the at-a-glance table."""
    return (
        compared.groupby(["ID", "Status"]).size().unstack(fill_value=0)
        .reindex(columns=["match", "differs", "expected_absence", "unexpected_absence", "unexpected_new_only"],
                 fill_value=0)
        .sort_index()
    )


def write_detail_workbook(compared: pd.DataFrame, path) -> None:
    """Writes the full comparison to `path`: a "Summary" sheet (one row per
    ID, from summarize()) and a "Details" sheet (every row not a plain match
    or expected absence — the console output only ever prints a sample of
    this)."""
    details = compared[~compared.Status.isin(["match", "expected_absence"])]
    with pd.ExcelWriter(path) as writer:
        summarize(compared).to_excel(writer, sheet_name="Summary")
        details.to_excel(writer, sheet_name="Details", index=False)


def main():
    prefix = sys.argv[1] if len(sys.argv) > 1 else ""
    sample_size = int(sys.argv[2]) if len(sys.argv) > 2 else 5

    original = load_all_original()
    consolidated = load_consolidated(extract.HISTORICAL_YEAR)
    if prefix:
        original = original[original.ID.str.startswith(prefix)]
        consolidated = consolidated[consolidated.ID.str.startswith(prefix)]

    compared = compare(original, consolidated, known_absent_from_consolidated())

    report_path = config.OUTPUT_DIR / f"Comparison_archived_vs_consolidated_{date.today().isoformat()}.xlsx"
    write_detail_workbook(compared, report_path)
    print(f"Full detail written to {report_path}\n")

    print(f"=== Per-ID summary ({extract.HISTORICAL_YEAR}"
          f"{f', ID starts with {prefix!r}' if prefix else ''}) ===")
    print(summarize(compared).to_string())

    interesting = compared[compared.Status.isin(["differs", "unexpected_absence", "unexpected_new_only"])]
    if interesting.empty:
        print("\nNothing to show in detail — no differences, no unexpected absences.")
        return

    print("\n=== Sample rows for anything not a plain match or expected absence ===")
    for id_ in sorted(interesting.ID.unique()):
        sub = interesting[interesting.ID == id_]
        print(f"\n--- {id_} ---")
        print(sub[["BU", "Month", "Original", "New", "Status"]].head(sample_size).to_string(index=False))


if __name__ == "__main__":
    main()
