"""Tests for scripts/compare_archived_vs_consolidated.py.

Real archive/consolidated files aren't available in CI (shared SharePoint
drive, confidential data) — everything here uses tiny synthetic fixtures,
following the same pattern as validate_against_originals's tests in
test_csr_calc_engine.py (_write_fixture_workbook + monkeypatch on config
paths).

known_absent_from_consolidated's own correctness against the real, current
pipeline state is already covered indirectly by
test_formulas_only_contains_additive_indicators and
test_contextual_values_covers_exactly_the_reclassified_indicators in
test_csr_calc_engine.py: if FORMULAS/CONTEXTUAL_VALUES ever drift, those
fail first. What's tested here is the derivation logic itself, with
monkeypatched inputs so it doesn't silently rely on today's exact indicator
set.
"""
import openpyxl
import pandas as pd
import pytest

import compare_archived_vs_consolidated as cmp
import config
import csr_calc_engine as engine
import extract_reference_and_data as extract


def _write_raw_workbook(path, sheets):
    """sheets: {month_name: {indicator_id: value}} -> a workbook shaped like
    the original per-BU raw files (ID in column B, value in column J)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for month, id_to_value in sheets.items():
        ws = wb.create_sheet(month)
        for row, (kpi_id, value) in enumerate(id_to_value.items(), start=2):
            ws.cell(row=row, column=2, value=kpi_id)
            ws.cell(row=row, column=10, value=value)
    wb.save(path)


@pytest.fixture
def single_bu(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "RAW_FILES", {"TestBU": "testbu.xlsx"})
    monkeypatch.setattr(config, "RAW_DATA_DIR", tmp_path)
    return tmp_path / "testbu.xlsx"


def test_read_original_indicators_reads_every_id_regardless_of_kind(single_bu):
    # Unlike extract_reference_and_data.build_raw_long (input-only), this
    # tool must also see calculated/excluded IDs -- that's the whole point,
    # comparing what used to be there against what the pipeline keeps today.
    _write_raw_workbook(single_bu, {"January": {"Env.1": 1000.0, "Wat.1": 5.0}})

    df = cmp.read_original_indicators("TestBU")

    assert set(df.ID) == {"Env.1", "Wat.1"}
    assert df[df.ID == "Wat.1"].iloc[0]["Original"] == 5.0


def test_load_consolidated_filters_to_requested_year(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    consolidated = pd.DataFrame([
        {"BU": "TestBU", "Year": 2026, "Month": "January", "ID": "Wat.1", "Value": 5.0},
        {"BU": "TestBU", "Year": 2027, "Month": "January", "ID": "Wat.1", "Value": 6.0},
    ])
    with pd.ExcelWriter(tmp_path / "Consolidated_results_CSR.xlsx") as writer:
        consolidated.to_excel(writer, sheet_name="Results", index=False)

    df = cmp.load_consolidated(2026)

    assert len(df) == 1
    assert df.iloc[0]["New"] == 5.0


def test_load_consolidated_raises_clear_error_when_file_missing(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    with pytest.raises(FileNotFoundError):
        cmp.load_consolidated(2026)


def test_known_absent_combines_excluded_and_uncomputed_calc_ids(monkeypatch):
    monkeypatch.setattr(extract, "EXCLUDED_IDS", {"Env.1"})
    monkeypatch.setattr(extract, "CALC_IDS", {"Ref.1", "Wat.2"})
    monkeypatch.setattr(engine, "FORMULAS", {"Ref.1": lambda v: 0})
    monkeypatch.setattr(engine, "CONTEXTUAL_VALUES", {})

    result = cmp.known_absent_from_consolidated()

    # Env.1: excluded outright. Wat.2: classified "calculated" but not
    # actually produced by FORMULAS/CONTEXTUAL_VALUES. Ref.1: produced by
    # FORMULAS, so NOT expected to be absent.
    assert result == {"Env.1", "Wat.2"}


def test_compare_classifies_every_status():
    original = pd.DataFrame([
        {"BU": "BU1", "Month": "January", "ID": "Wat.1", "Original": 5.0},   # match
        {"BU": "BU1", "Month": "January", "ID": "Saf.1", "Original": 10.0},  # differs
        {"BU": "BU1", "Month": "January", "ID": "Env.1", "Original": 999.0}, # expected_absence
        {"BU": "BU1", "Month": "January", "ID": "Was.3", "Original": 1.0},   # unexpected_absence
    ])
    consolidated = pd.DataFrame([
        {"BU": "BU1", "Month": "January", "ID": "Wat.1", "New": 5.0},
        {"BU": "BU1", "Month": "January", "ID": "Saf.1", "New": 12.0},
        {"BU": "BU1", "Month": "January", "ID": "Ene.9", "New": 42.0},  # unexpected_new_only
    ])

    result = cmp.compare(original, consolidated, known_absent={"Env.1"})
    status = result.set_index("ID")["Status"]

    assert status["Wat.1"] == "match"
    assert status["Saf.1"] == "differs"
    assert status["Env.1"] == "expected_absence"
    assert status["Was.3"] == "unexpected_absence"
    assert status["Ene.9"] == "unexpected_new_only"


def test_compare_treats_none_and_nan_as_a_match():
    # A blank original cell (None) vs. an unfilled consolidated value (NaN,
    # after the outer merge) must not be flagged as "differs".
    original = pd.DataFrame([{"BU": "BU1", "Month": "January", "ID": "Wat.1", "Original": None}])
    consolidated = pd.DataFrame([{"BU": "BU1", "Month": "January", "ID": "Wat.1", "New": None}])

    result = cmp.compare(original, consolidated)

    assert result.iloc[0]["Status"] == "match"


def test_write_detail_workbook_excludes_matches_and_expected_absences_from_details(tmp_path):
    compared = pd.DataFrame([
        {"BU": "BU1", "Month": "January", "ID": "Wat.1", "Original": 5.0, "New": 5.0, "Status": "match"},
        {"BU": "BU1", "Month": "January", "ID": "Env.1", "Original": 999.0, "New": None, "Status": "expected_absence"},
        {"BU": "BU1", "Month": "January", "ID": "Saf.1", "Original": 10.0, "New": 12.0, "Status": "differs"},
    ])
    path = tmp_path / "report.xlsx"

    cmp.write_detail_workbook(compared, path)

    summary = pd.read_excel(path, sheet_name="Summary", index_col=0)
    assert summary.loc["Wat.1", "match"] == 1
    details = pd.read_excel(path, sheet_name="Details")
    assert list(details.ID) == ["Saf.1"]


def test_summarize_counts_rows_per_id_and_status():
    compared = pd.DataFrame([
        {"ID": "Wat.1", "Status": "match"},
        {"ID": "Wat.1", "Status": "match"},
        {"ID": "Wat.1", "Status": "differs"},
        {"ID": "Saf.1", "Status": "expected_absence"},
    ])

    summary = cmp.summarize(compared)

    assert summary.loc["Wat.1", "match"] == 2
    assert summary.loc["Wat.1", "differs"] == 1
    assert summary.loc["Saf.1", "expected_absence"] == 1
    assert summary.loc["Saf.1", "match"] == 0
